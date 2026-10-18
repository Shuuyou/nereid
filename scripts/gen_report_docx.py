#!/usr/bin/env python3
"""技术报告 markdown → docx：解析标题/表格/代码块/列表，套规范样式。
用法: python3 scripts/gen_report_docx.py docs/08-技术报告.md docs/捞针·Nereid-技术报告.docx
PDF: libreoffice --headless --convert-to pdf <docx>
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def _cn_font(run, name="宋体", size=None, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def _style(doc):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for lvl, sz in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        s = doc.styles[lvl]
        s.font.name = "Times New Roman"
        s.font.size = Pt(sz)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")


def _inline(p, text):
    """粗体/行内代码简易解析"""
    for i, seg in enumerate(re.split(r"(\*\*.+?\*\*|`.+?`)", text)):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            r = p.add_run(seg[2:-2])
            r.bold = True
            _cn_font(r)
        elif seg.startswith("`") and seg.endswith("`"):
            r = p.add_run(seg[1:-1])
            _cn_font(r, "Consolas", 10)
        else:
            _cn_font(p.add_run(seg))


def build(src: str, dst: str) -> None:
    doc = Document()
    _style(doc)
    lines = open(src, encoding="utf-8").read().splitlines()
    i = 0
    in_code = False
    code_buf: list[str] = []
    table_buf: list[str] = []

    def flush_table():
        nonlocal table_buf
        if not table_buf:
            return
        rows = [r for r in table_buf if not re.match(r"^\s*\|?[\s:\-|]+\|?\s*$", r)]
        cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        if cells:
            t = doc.add_table(rows=len(cells), cols=len(cells[0]))
            t.style = "Table Grid"
            for ri, row in enumerate(cells):
                for ci, cell in enumerate(row[:len(cells[0])]):
                    p = t.cell(ri, ci).paragraphs[0]
                    _inline(p, cell)
                    if ri == 0:
                        for r in p.runs:
                            r.bold = True
        table_buf = []

    while i < len(lines):
        ln = lines[i]
        if ln.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                for cl in code_buf:
                    r = p.add_run(cl + "\n")
                    _cn_font(r, "Consolas", 9)
                p.paragraph_format.left_indent = Pt(18)
                code_buf = []
                in_code = False
            else:
                flush_table()
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(ln)
            i += 1
            continue
        if "|" in ln and ln.strip().startswith("|"):
            table_buf.append(ln)
            i += 1
            continue
        flush_table()
        s = ln.strip()
        if not s:
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            lvl = min(len(m.group(1)), 3)
            doc.add_heading(m.group(2).strip(), level=lvl)
        elif re.match(r"^[-*]\s+", s):
            p = doc.add_paragraph(style="List Bullet")
            _inline(p, re.sub(r"^[-*]\s+", "", s))
        elif re.match(r"^\d+[.)]\s+", s):
            p = doc.add_paragraph(style="List Number")
            _inline(p, re.sub(r"^\d+[.)]\s+", "", s))
        elif s.startswith(">"):
            p = doc.add_paragraph()
            _inline(p, s.lstrip("> "))
            p.paragraph_format.left_indent = Pt(18)
        elif s.startswith("---"):
            pass
        else:
            p = doc.add_paragraph()
            _inline(p, s)
        i += 1
    flush_table()
    doc.save(dst)
    print("docx:", dst)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
